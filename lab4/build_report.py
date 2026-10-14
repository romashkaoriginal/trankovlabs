from __future__ import annotations

from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
REPORT = ROOT / "ЛР4_Траньков.docx"
DIAGRAM = ROOT / "uml_auth.png"
RESET_DIAGRAM = ROOT / "uml_recovery.png"
STRUCTURE_DIAGRAM = ROOT / "subsystems.png"
IDEF0_EXPORTS = ROOT / "diagrams"
FONT = Path(r"C:\Windows\Fonts\times.ttf")
FONT_BOLD = Path(r"C:\Windows\Fonts\timesbd.ttf")


def draw_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], width: int = 5) -> None:
    draw.line((start, end), fill="black", width=width)
    x, y = end
    if start[0] == end[0]:
        direction = 1 if end[1] > start[1] else -1
        draw.polygon([(x, y), (x - 12, y - direction * 19), (x + 12, y - direction * 19)], fill="black")
    else:
        direction = 1 if end[0] > start[0] else -1
        draw.polygon([(x, y), (x - direction * 19, y - 12), (x - direction * 19, y + 12)], fill="black")


def make_diagram() -> None:
    img = Image.new("RGB", (1600, 1870), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype(str(FONT), 36)
    small = ImageFont.truetype(str(FONT), 32)
    bold = ImageFont.truetype(str(FONT_BOLD), 38)

    def centered(lines: str, box: tuple[int, int, int, int], face=font) -> None:
        x1, y1, x2, y2 = box
        parts = lines.split("\n")
        total = len(parts) * 48
        y = (y1 + y2 - total) // 2
        for line in parts:
            bounds = draw.textbbox((0, 0), line, font=face)
            draw.text(((x1 + x2 - (bounds[2] - bounds[0])) // 2, y), line, fill="black", font=face)
            y += 48

    def step(box: tuple[int, int, int, int], label: str) -> None:
        draw.rounded_rectangle(box, radius=18, outline="black", width=5)
        centered(label, box)

    def decision(y: int, label: str) -> None:
        draw.polygon([(800, y - 82), (1040, y), (800, y + 82), (560, y)], outline="black", fill="white", width=5)
        centered(label, (625, y - 70, 975, y + 70), small)

    draw.ellipse((777, 35, 823, 81), fill="black")
    draw_arrow(draw, (800, 81), (800, 150))
    step((445, 150, 1155, 270), "Пользователь вводит\nлогин и пароль")
    draw_arrow(draw, (800, 270), (800, 340))
    step((445, 340, 1155, 470), "Сервер проверяет учётную\nзапись и пароль")
    draw_arrow(draw, (800, 470), (800, 555))
    decision(640, "Данные\nверны?")

    draw_arrow(draw, (1040, 640), (1130, 640))
    draw.text((1050, 584), "нет", fill="black", font=small)
    step((1130, 580, 1550, 700), "Показать ошибку\nвхода")
    draw.line((1340, 700, 1340, 770), fill="black", width=5)
    draw.ellipse((1322, 770, 1358, 806), outline="black", width=5)
    draw.ellipse((1329, 777, 1351, 799), fill="black")

    draw_arrow(draw, (800, 722), (800, 790))
    draw.text((830, 737), "да", fill="black", font=small)
    step((445, 790, 1155, 910), "Создать сеанс и\nопределить роль")
    draw_arrow(draw, (800, 910), (800, 980))
    step((445, 980, 1155, 1090), "Пользователь запрашивает действие")
    draw_arrow(draw, (800, 1090), (800, 1150))
    decision(1235, "Есть право?")

    draw_arrow(draw, (1040, 1235), (1130, 1235))
    draw.text((1050, 1178), "нет", fill="black", font=small)
    step((1130, 1175, 1550, 1295), "Отказать\nв доступе")
    draw.line((1340, 1295, 1340, 1365), fill="black", width=5)
    draw.ellipse((1322, 1365, 1358, 1401), outline="black", width=5)
    draw.ellipse((1329, 1372, 1351, 1394), fill="black")

    draw_arrow(draw, (800, 1317), (800, 1390))
    draw.text((830, 1335), "да", fill="black", font=small)
    step((445, 1390, 1155, 1500), "Выполнить действие")
    draw_arrow(draw, (800, 1500), (800, 1570))
    step((400, 1570, 1200, 1720), "Через 10 минут бездействия\nсервер завершает сеанс")
    draw_arrow(draw, (800, 1720), (800, 1800))
    draw.ellipse((776, 1800, 824, 1848), outline="black", width=5)
    draw.ellipse((785, 1809, 815, 1839), fill="black")

    img.save(DIAGRAM, dpi=(240, 240))


def set_cell_border(cell, size=4) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        item = borders.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), str(size))
        item.set(qn("w:color"), "000000")


def set_padding(cell, twips=40) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side in ("top", "left", "bottom", "right"):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(twips))
        node.set(qn("w:type"), "dxa")


def field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)


def set_header(section, doc: Document, width_cm: float) -> None:
    section.header.is_linked_to_previous = False
    p = section.header.paragraphs[0]
    p.clear()
    p.style = doc.styles["Normal"]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(width_cm), WD_TAB_ALIGNMENT.RIGHT)
    for value in ("СИ-41", "\tТраньков Р.С."):
        run = p.add_run(value)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)


def set_report_header(section, doc: Document) -> None:
    section.header.is_linked_to_previous = False
    section.header_distance = Cm(0.45)
    p = section.header.paragraphs[0]
    p.clear()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 0.1
    table = section.header.add_table(rows=1, cols=4, width=Cm(17))
    table.autofit = False
    widths = [3.2, 8.0, 2.9, 2.9]
    cells = table.rows[0].cells
    labels = [
        "СИ-41\nТраньков Р.С.",
        "Технические требования\nк библиотечной информационной системе",
        "Номер текущего листа\n",
        "Всего листов\n",
    ]
    for i, (cell, label) in enumerate(zip(cells, labels)):
        cell.width = Cm(widths[i])
        cell.text = label
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell)
        set_padding(cell, 35)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(8)
    for i, instruction in [(2, "PAGE"), (3, "NUMPAGES")]:
        field(cells[i].paragraphs[0], instruction)
        cells[i].paragraphs[0].runs[-1].font.size = Pt(8)
    section.footer.is_linked_to_previous = False
    section.footer.paragraphs[0].clear()


def blank_line(doc: Document):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    return p


def normal_paragraph(doc: Document, text: str, *, indent=True):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1) if indent else Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.add_run(text)
    return p


def intro_item(doc: Document, label: str, text: str):
    p = normal_paragraph(doc, "")
    p.add_run(label).bold = True
    p.add_run(" " + text)
    return p


def heading(doc: Document, number: str, title: str, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(f"{number}\u00a0{title}")
    blank_line(doc)
    return p


def requirement(doc: Document, req_id: str, statement: str, check: str):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(req_id + ". ")
    r.bold = True
    p.add_run(statement + " Проверка: " + check)


def table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], widths: list[float]):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(12)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, (label, width) in enumerate(zip(headers, widths)):
        t.columns[i].width = Cm(width)
        c = t.rows[0].cells[i]
        c.width = Cm(width)
        c.text = label
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in c.paragraphs[0].runs:
            run.bold = True
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_row = t.rows[0]._tr
    tr_pr = header_row.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for data in rows:
        cells = t.add_row().cells
        for i, value in enumerate(data):
            cells[i].width = Cm(widths[i])
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tr_pr = t.rows[-1]._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)
    for row_index, row in enumerate(t.rows):
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for col_index, cell in enumerate(row.cells):
            set_cell_border(cell)
            set_padding(cell)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 or widths[col_index] <= 2.5 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10 if row_index == 0 else 12)
                    run.bold = row_index == 0
    return t


REQUIREMENTS = [
    ("S-01", "Система состоит из подсистем доступа, каталога и фонда, выдачи и возврата, уведомлений и отчётов.", "В интерфейсе доступны функции каждой подсистемы согласно роли пользователя.", "A0, M2"),
    ("S-02", "Сотрудник работает с системой через браузер. Библиотечные операции не требуют установки отдельного клиентского приложения.", "Операции S-01 выполняются в браузере на рабочем месте библиотекаря.", "M1, M2"),
    ("A-01", "Система запрашивает логин и пароль перед доступом к данным читателей и операциям фонда.", "Без входа защищённые страницы и запросы не возвращают данные.", "ЛР IDEF0: M2; уточнение доступа"),
    ("A-02", "При создании и восстановлении пароля интерфейс проверяет длину не менее 8 символов, наличие цифры, строчной и прописной буквы. Сервер повторяет ту же проверку.", "Отклоняются короткий пароль, пароль без цифры или без одного из регистров, в том числе при прямом запросе к серверу.", "ЛР IDEF0: M2; уточнение доступа"),
    ("A-03", "Система хранит пароль как стойкий односторонний хеш с индивидуальной солью; исходный пароль и его обратимо зашифрованная копия не сохраняются.", "В базе нет исходного пароля; два одинаковых пароля имеют разные сохранённые значения.", "ЛР IDEF0: M2; уточнение доступа"),
    ("A-04", "После входа система определяет роль: библиотекарь или администратор. Права проверяются сервером перед каждой защищённой операцией.", "Библиотекарь не может управлять учётными записями, в том числе прямым запросом к серверу.", "ЛР IDEF0: M1, M2; уточнение доступа"),
    ("A-05", "После 10 минут бездействия сервер прекращает сеанс. Следующий защищённый запрос требует повторного входа.", "Через 10 минут без запросов прежний идентификатор сеанса не даёт доступа.", "ЛР IDEF0: M2; уточнение доступа"),
    ("A-06", "Пользователь может запросить сброс пароля. Ссылка или код восстановления действует один раз в течение 15 минут; после смены пароля прежние сеансы завершаются.", "Повторное использование кода и использование спустя 15 минут отклоняются.", "ЛР IDEF0: M2; уточнение доступа"),
    ("A-07", "Система регистрирует успешные и неуспешные попытки входа, сброс пароля и отказ в доступе без записи паролей в журнал.", "В журнале есть тип события, время и учётная запись; значения пароля нет.", "ЛР IDEF0: M2; уточнение доступа"),
    ("F-01", "Библиотекарь ищет издание в каталоге и видит доступность конкретных экземпляров.", "Поиск возвращает издание и статус каждого найденного экземпляра.", "A1, C2, M3"),
    ("F-02", "До выдачи система проверяет учётную запись читателя, доступность экземпляра, лимиты и сведения о просрочке.", "При нарушении любого условия выдача блокируется с указанием причины.", "A1, I1, C3, W3"),
    ("F-03", "При выдаче система создаёт запись с читателем, экземпляром и сроком возврата, затем меняет статус экземпляра на «выдан» в одной операции.", "После успешной выдачи запись и статус совпадают; при ошибке оба изменения отменяются.", "A2, W1, W2, O1, O2"),
    ("F-04", "При возврате система закрывает открытую запись выдачи и переводит экземпляр в доступный статус.", "После возврата нет открытой выдачи данного экземпляра, а статус фонда обновлён.", "A3, I2, O2"),
    ("F-05", "Система отмечает просроченные выдачи и формирует уведомление читателю; сведения о просрочке учитываются при новой заявке.", "Просроченная запись попадает в уведомления и влияет на проверку F-02.", "A3, A4, W3, O3"),
    ("F-06", "Библиотекарь формирует отчёт по выдачам, возвратам и просрочкам за выбранный период.", "Итоги отчёта совпадают с записями операций за тот же период.", "A4, O4"),
    ("Q-01", "Система не допускает одновременной открытой выдачи одного экземпляра двум читателям.", "Два параллельных запроса на выдачу оставляют не более одной открытой записи.", "A2, O2"),
    ("Q-02", "Данные операций и учётных записей резервируются; восстановление проверяется на отдельной копии.", "Из резервной копии восстанавливаются записи выдачи, возврата и учётные записи.", "M2; проектное требование"),
    ("I-01", "Логин, пароль и идентификатор сеанса передаются по защищённому соединению; сервер отклоняет передачу учётных данных по открытому HTTP.", "Запрос с учётными данными по HTTP не обрабатывается.", "ЛР IDEF0: M2; уточнение доступа"),
    ("I-02", "Библиотекарь видит данные читателя только для обслуживания, а администратор управляет доступом без просмотра исходных паролей.", "Проверка ролей не открывает пароли и не даёт доступа к операциям вне назначенной роли.", "M1, M2"),
    ("D-01", "Система хранит читателя, экземпляр, запись выдачи и возврата, учётную запись сотрудника и события доступа как отдельные связанные записи.", "По идентификаторам можно восстановить историю операций экземпляра и читателя.", "I1, I2, C2, W2, M2"),
    ("DOC-01", "Комплект документации включает технические требования, описание ролей, руководство библиотекаря, инструкцию администратора и UML-диаграмму входа.", "Все перечисленные документы представлены при приёмке.", "Состав документации проекта"),
]

# Значения ниже являются проектными допущениями для проверяемого учебного ТТ.
REVISED_REQUIREMENTS = {
    "S-02": ("Сотрудник выполняет библиотечные операции в Chrome 128+, Edge 128+ или Firefox 128+ без установки отдельного клиента.", "Поиск, выдача и возврат проходят в каждом из трёх браузеров указанных версий.", "M1, M2; проектное допущение"),
    "A-02": ("При создании и восстановлении пароля интерфейс проверяет длину от 8 до 128 символов, наличие цифры, строчной и прописной буквы. Допускаются буквы Unicode, цифры и печатные знаки; сервер повторяет проверку.", "Пароль из 7 и 129 символов отклоняется; подходящие пароли из 8 и 128 символов принимаются. Прямой запрос не обходит правило.", "M2; требование задания и проектное уточнение"),
    "A-03": ("Пароли хранятся только как хеш Argon2id с индивидуальной случайной солью не менее 16 байт; параметры: m=19456 КиБ, t=2, p=1. Исходный пароль не сохраняется.", "В хранилище нет исходного пароля; два одинаковых пароля дают разные записи, параметры Argon2id совпадают с заданными.", "M2; OWASP [1]"),
    "A-04": ("После входа сервер назначает роль и проверяет её перед каждой операцией. Библиотекарь обслуживает читателей и фонд; администратор управляет учётными записями и настройками, но не оформляет выдачу.", "Запросы вне разрешённой роли отклоняются, в том числе при прямом обращении к серверу.", "M1, M2; проектное расширение роли"),
    "A-05": ("Сеанс завершается через 10 минут после последнего авторизованного запроса к серверу. За минуту до завершения интерфейс предупреждает пользователя и сохраняет несохранённую форму локально до повторного входа.", "Запрос в 9:59 продлевает сеанс; после 10:01 без запросов прежний идентификатор не даёт доступа, введённая форма доступна после повторного входа.", "M2; требование задания"),
    "A-06": ("Восстановление логина и пароля выполняется через адрес электронной почты, заранее подтверждённый для учётной записи. Одноразовый код действует 15 минут, допускает не более пяти попыток; ответ не раскрывает наличие учётной записи. После смены пароля прежние сеансы закрываются.", "Проверяются одноразовость, срок 15 минут, шестая попытка, одинаковый ответ для известного и неизвестного адреса и завершение старых сеансов.", "M2; проектное расширение доступа"),
    "A-07": ("Система записывает вход, отказ, сброс пароля и завершение сеанса без паролей и кодов; журнал хранится 90 дней.", "Запись содержит время, тип события и идентификатор учётной записи; секретов нет, записи старше 90 дней удаляются.", "M2; проектное требование"),
    "F-02": ("Перед выдачей система проверяет активность читателя, доступность экземпляра, отсутствие просрочки и лимит не более пяти открытых выдач; срок обычной выдачи составляет 14 календарных дней.", "Шестая открытая выдача, просрочка или недоступный экземпляр блокируют операцию с причиной; срок возврата равен дате выдачи плюс 14 дней.", "A1, I1, C1, C2, C3, W1"),
    "F-03": ("При принятом решении по заявке система создаёт запись выдачи с читателем, экземпляром и сроком возврата, меняет статус экземпляра на «выдан» и передаёт экземпляр читателю в одной операции.", "После успешной выдачи запись и статус совпадают; при ошибке оба изменения отменяются, экземпляр не отмечен как переданный.", "A2, W1, W2, O1"),
    "F-04": ("При возврате библиотекарь отмечает состояние экземпляра; система закрывает выдачу, фиксирует повреждение при наличии и обновляет статус фонда.", "После возврата нет открытой выдачи, состояние и статус экземпляра записаны, повреждённый экземпляр не доступен для новой выдачи.", "A3, I2, W2, O2"),
    "F-05": ("Просроченной считается выдача после даты возврата при отсутствии факта возврата. Система передаёт читателю уведомление по электронной почте и фиксирует результат доставки.", "На следующий день после срока возврата формируется уведомление; ошибка почтового сервера фиксируется для повторной отправки.", "A3, A4, W3, O3; почтовый шлюз"),
    "F-06": ("Библиотекарь выгружает отчёт в PDF и CSV за выбранный период; каждая строка содержит читателя, экземпляр, даты выдачи и возврата, статус просрочки.", "Состав строк и итоги совпадают с записями периода; PDF и CSV содержат названные поля.", "A4, O4"),
    "Q-02": ("Полная резервная копия данных создаётся не реже одного раза в сутки и хранится 30 дней; восстановление на отдельном стенде занимает не более четырёх часов.", "Проверяются дата последней копии, наличие копий за 30 дней и восстановление в пределах четырёх часов.", "M2; проектное допущение"),
    "I-01": ("Учётные данные и идентификатор сеанса передаются только через TLS 1.2 или выше; запросы по открытому HTTP не обрабатываются.", "Подключение по TLS 1.1 и HTTP отклоняется, по TLS 1.2 проходит.", "M2; проектное требование"),
    "I-02": ("Библиотекарь получает только поля читателя, нужные для обслуживания; администратор не получает исходные пароли и хеши через интерфейс или API.", "Ответы API каждой роли не содержат лишних полей и хешей паролей.", "M1, M2; проектное требование"),
}
REQUIREMENTS = [(rid, *REVISED_REQUIREMENTS[rid]) if rid in REVISED_REQUIREMENTS else (rid, text, check, source)
                for rid, text, check, source in REQUIREMENTS]
REQUIREMENTS.extend([
    ("A-08", "После пяти неуспешных попыток входа учётная запись блокируется на 15 минут; время блокировки не продлевается повторными запросами.", "Пятая ошибка включает блокировку, верный пароль во время блокировки отклоняется, через 15 минут вход возможен.", "M2; проектное уточнение доступа"),
    ("A-09", "Ошибка входа не сообщает, какой из реквизитов неверен и существует ли учётная запись.", "Для неверного логина и неверного пароля показывается одинаковый ответ.", "M2; проектное уточнение доступа"),
    ("F-07", "Библиотекарь создаёт карточку читателя и запись каталога и может обновлять контактные сведения читателя, но не права сотрудников.", "Созданные записи доступны при поиске; запрос библиотекаря на изменение роли сотрудника отклоняется.", "C2, M1, M2; проектное дополнение"),
    ("Q-03", "При 20 одновременных сотрудниках 95 % запросов поиска выполняются не более чем за 3 секунды.", "Нагрузочное испытание с 20 сеансами подтверждает показатель на тестовом каталоге.", "M2; проектное допущение"),
    ("Q-04", "Доступность БИС составляет не менее 99 % за календарный месяц без учёта заранее объявленного обслуживания.", "По журналу мониторинга рассчитывается месячная доля доступности.", "M2; проектное допущение"),
    ("D-02", "Запрос читателя содержит идентификатор читателя и экземпляра; отчёт содержит период, перечень операций и итоговые количества выдач, возвратов и просрочек.", "Запрос без обязательного поля отклоняется; отчёт содержит все перечисленные поля.", "I1, O4; проектное уточнение ICOM"),
    ("INT-01", "Уведомления отправляются через почтовый сервер по SMTP с TLS; при недоступности сервера сообщение ставится в очередь повторной отправки.", "Успешная и неуспешная отправки фиксируются, недоставленное сообщение повторяется после восстановления связи.", "O3; проектное смежное взаимодействие"),
    ("PAT-01", "Перед передачей БИС заказчику состав компонентов и условия их лицензирования документируются; компонент без подтверждённого права использования не включается в поставку.", "В перечне указаны компоненты и лицензии, неподтверждённый компонент отсутствует в поставке.", "Документация проекта"),
])
ACCESS_SOURCES = {
    "A-01": "Б.1: вход; M2 (расширение)",
    "A-02": "Задание ЛР4; Б.2: пароль",
    "A-03": "Задание ЛР4; OWASP [1]",
    "A-04": "Б.1: проверка роли; M1, M2",
    "A-05": "Задание ЛР4; Б.1: таймер",
    "A-06": "Задание ЛР4; Б.2: сброс",
    "A-07": "Б.1: журнал; M2",
    "A-08": "Б.1: защита от подбора",
    "A-09": "Б.1: отказ во входе",
    "I-01": "Задание ЛР4; M2",
    "I-02": "A-04; M1, M2",
    "Q-01": "A2, W2; проектное требование",
}
REQUIREMENTS = [(rid, statement, check, ACCESS_SOURCES.get(rid, source))
                for rid, statement, check, source in REQUIREMENTS]

TRACE_HINTS = {
    "S-01": "Состав подсистем БИС.",
    "S-02": "Рабочее место библиотекаря.",
    "A-01": "Вход по логину и паролю.",
    "A-02": "Правила нового пароля.",
    "A-03": "Защищённое хранение паролей.",
    "A-04": "Разграничение прав сотрудников.",
    "A-05": "Завершение сеанса после бездействия.",
    "A-06": "Восстановление доступа.",
    "A-07": "Учёт событий доступа.",
    "F-01": "Поиск и доступность экземпляров.",
    "F-02": "Проверка заявки перед выдачей.",
    "F-03": "Оформление выдачи.",
    "F-04": "Оформление возврата.",
    "F-05": "Просрочка и уведомление читателя.",
    "F-06": "Отчёт по операциям.",
    "Q-01": "Запрет двойной выдачи.",
    "Q-02": "Резервирование данных БИС.",
    "I-01": "Защита данных при передаче.",
    "I-02": "Доступ к данным по роли.",
    "D-01": "Состав связанных записей.",
    "DOC-01": "Комплект документов при приёмке.",
    "A-08": "Ограничение неудачных попыток входа.",
    "A-09": "Одинаковый ответ при ошибке входа.",
    "F-07": "Создание карточек читателей и фонда.",
    "Q-03": "Проверка времени отклика при нагрузке.",
    "Q-04": "Месячная доступность системы.",
    "D-02": "Обязательные поля запроса и отчёта.",
    "INT-01": "Передача уведомлений почтовому серверу.",
    "PAT-01": "Проверка права использования компонентов.",
}


def add_requirements(doc: Document, ids: list[str]) -> None:
    for req_id, statement, check, _ in REQUIREMENTS:
        if req_id in ids:
            requirement(doc, req_id, statement, check)


def make_revised_diagrams() -> None:
    font = ImageFont.truetype(str(FONT), 35)
    small = ImageFont.truetype(str(FONT), 28)
    bold = ImageFont.truetype(str(FONT_BOLD), 39)

    def base(title: str, rows: list[tuple[str, str]], out: Path,
             branches: list[tuple[int, int, str, str]] | None = None,
             skip_vertical: set[int] | None = None) -> None:
        width, row_h = 1700, 168
        height = 220 + row_h * len(rows) + 150
        img = Image.new("RGB", (width, height), "white")
        d = ImageDraw.Draw(img)
        d.text((70, 40), title, font=bold, fill="black")
        d.line((70, 102, width - 70, 102), fill="black", width=3)
        x1, x2 = 260, 1420
        for i, (kind, label) in enumerate(rows):
            y1 = 155 + i * row_h
            y2 = y1 + 118
            if kind == "condition":
                d.polygon([(840, y1), (x2, (y1 + y2)//2), (840, y2), (x1, (y1 + y2)//2)], outline="black", fill="white", width=4)
            else:
                d.rounded_rectangle((x1, y1, x2, y2), radius=18, outline="black", width=4)
            lines = label.split("\n")
            line_h = 41
            ty = (y1 + y2 - len(lines) * line_h)//2
            for line in lines:
                bb = d.textbbox((0, 0), line, font=font)
                d.text(((width-(bb[2]-bb[0]))//2, ty), line, font=font, fill="black")
                ty += line_h
            if i + 1 < len(rows) and i not in (skip_vertical or set()):
                draw_arrow(d, (840, y2), (840, y2 + 46), 4)
        for source, target, side, label in branches or []:
            x = 1520 if side == "right" else 170
            sx = x2 if side == "right" else x1
            sy = 155 + source * row_h + 59
            ty = 155 + target * row_h + 59
            d.line([(sx, sy), (x, sy), (x, ty), (sx, ty)], fill="black", width=4)
            draw_arrow(d, (x, ty), (sx, ty), 4)
            d.text((x + 12 if side == "right" else 65, min(sy, ty) + 10), label, font=small, fill="black")
        if out != STRUCTURE_DIAGRAM:
            d.ellipse((825, 112, 855, 142), fill="black")
            draw_arrow(d, (840, 142), (840, 155), 3)
            last_y = 155 + (len(rows)-1) * row_h + 118
            draw_arrow(d, (840, last_y), (840, last_y + 64), 3)
            d.ellipse((815, last_y + 67, 865, last_y + 117), outline="black", width=4)
            d.ellipse((824, last_y + 76, 856, last_y + 108), fill="black")
        img.save(out, dpi=(240, 240))

    base("Б.1. Вход и работа в системе", [
        ("action", "Сотрудник открывает БИС"),
        ("action", "Ввести логин и пароль"),
        ("action", "Сервер проверяет пароль и блокировку;\nсобытие записывается в журнал"),
        ("condition", "[Данные верны и запись не заблокирована?]"),
        ("action", "[нет] Единая ошибка; повторить ввод.\nПосле 5 ошибок — блокировка на 15 минут"),
        ("action", "[да] Создать сеанс и определить роль"),
        ("condition", "[Есть новый запрос?]"),
        ("action", "[да] Проверить право на операцию;\nразрешить либо отказать и записать событие"),
        ("action", "Вернуться к ожиданию следующего запроса;\nзапрос продлевает сеанс"),
        ("action", "[нет] Выход пользователя или 10 минут\nбез запросов → завершить сеанс"),
    ], DIAGRAM, [(3, 5, "right", "[да]"), (4, 1, "left", "[повтор]"),
                 (6, 9, "right", "[нет]"), (8, 6, "left", "[ещё запрос]")], {4, 8})
    base("Б.2. Создание пароля и восстановление доступа", [
        ("action", "Создать пароль или восстановить доступ"),
        ("condition", "[Запрошено восстановление?]"),
        ("action", "[да] Ввести подтверждённый e-mail;\nответ одинаков для любого адреса"),
        ("action", "Отправить одноразовый код на 15 минут;\nдопустить не более 5 попыток"),
        ("condition", "[Код действителен?]"),
        ("action", "[нет] Отказать; предложить новый запрос"),
        ("action", "[да / создание] Ввести новый пароль"),
        ("action", "Интерфейс проверяет 8–128 символов,\nцифру, строчную и прописную букву"),
        ("condition", "[Пароль соответствует правилу?]"),
        ("action", "[нет] Показать правило и повторить ввод"),
        ("action", "[да] Сервер повторяет проверку, сохраняет\nArgon2id-хеш и завершает старые сеансы"),
    ], RESET_DIAGRAM, [(1, 6, "left", "[нет]"), (4, 6, "right", "[да]"),
                       (5, 2, "right", "[повтор]"), (8, 10, "right", "[да]"),
                       (9, 6, "left", "[повтор]")], {5, 9})
    structure = Image.new("RGB", (1700, 900), "white")
    sd = ImageDraw.Draw(structure)
    sd.rectangle((420, 45, 1280, 155), outline="black", width=4)
    sd.text((493, 78), "Библиотечная информационная система", font=font, fill="black")
    sd.line((850, 155, 850, 765), fill="black", width=4)
    modules = [
        ("Доступ и роли", "Каталог и читатели"),
        ("Выдача и возврат", "Уведомления"),
        ("Отчёты", "Резервные копии"),
    ]
    for row, pair in enumerate(modules):
        y = 265 + row * 220
        for left, right, label in [(50, 745, pair[0]), (955, 1650, pair[1])]:
            sd.rectangle((left, y, right, y + 115), outline="black", width=4)
            bb = sd.textbbox((0, 0), label, font=font)
            sd.text(((left + right - (bb[2]-bb[0]))//2, y + 39), label, font=font, fill="black")
        sd.line((745, y + 58, 955, y + 58), fill="black", width=4)
    structure.save(STRUCTURE_DIAGRAM, dpi=(240, 240))


def build_document() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(3)
    sec.right_margin = Cm(1)
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2.5)
    sec.header_distance = Cm(0.6)
    sec.footer_distance = Cm(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(1)
    normal.paragraph_format.widow_control = True
    def use_times(style) -> None:
        rpr = style._element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{attr}"), None)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attr}"), "Times New Roman")

    use_times(normal)
    for name in ("Heading 1", "Heading 2"):
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(14)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.first_line_indent = Cm(0)
        st.paragraph_format.line_spacing = 1.0
        st.paragraph_format.keep_with_next = True
        use_times(st)
    title_style = doc.styles["Title"]
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(14)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_before = Pt(0)
    title_style.paragraph_format.space_after = Pt(0)
    title_style.paragraph_format.first_line_indent = Cm(0)
    title_style.paragraph_format.line_spacing = 1.0
    title_style.paragraph_format.keep_with_next = True
    use_times(title_style)
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    settings = doc.settings._element
    auto_hyphenation = settings.find(qn("w:autoHyphenation"))
    if auto_hyphenation is None:
        auto_hyphenation = OxmlElement("w:autoHyphenation")
        settings.append(auto_hyphenation)
    auto_hyphenation.set(qn("w:val"), "1")

    set_header(sec, doc, 17)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Cm(0)
    footer.paragraph_format.space_after = Pt(0)
    field(footer, "PAGE")

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("ЛАБОРАТОРНАЯ РАБОТА №\u00a04")
    p.add_run().add_break()
    p.add_run("РАЗРАБОТКА ТЕХНИЧЕСКИХ ТРЕБОВАНИЙ")
    blank_line(doc)
    intro_item(doc, "Цель:", "ознакомиться с порядком разработки технических требований к автоматизированной системе.")
    blank_line(doc)
    intro_item(doc, "Тема проекта:", "технические требования к библиотечной информационной системе учёта выдачи и возврата изданий на основе IDEF0-модели варианта 18.")
    blank_line(doc)
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Cm(1)
    p.add_run("Задачи работы:").bold = True
    for task in (
        "Выделить из IDEF0-модели функции, входы, выходы, нормативную базу, роли и смежные системы.",
        "Сформулировать проверяемые требования и присвоить каждому идентификатор.",
        "Описать вход в систему, права пользователей и восстановление доступа.",
        "Показать связь требований с моделью, подготовить UML-диаграмму и после утверждения разместить ТТ в репозитории проекта.",
    ):
        normal_paragraph(doc, "– " + task)
    blank_line(doc)
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Cm(1)
    p.add_run("Технические требования:").bold = True
    normal_paragraph(doc, "Задание требует входа по логину и паролю, проверки нового пароля на стороне клиента, хранения пароля «в зашифрованном виде», завершения сеанса после 10 минут бездействия и восстановления учётной записи. Структура ТТ задана с отсылкой к СТП 09150.11.118-2009. Проектное уточнение: сервер повторяет проверку пароля, а для хранения применяется Argon2id вместо обратимого шифрования. Это уточнение требует согласования с преподавателем.")
    blank_line(doc)
    intro_item(doc, "Общие сведения:", "технические требования описывают назначение БИС, её функции, ограничения и порядок проверки. Исходной моделью служит лабораторная работа IDEF0; связь с ней приведена в приложении В.")

    cover_sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    cover_sec.header.is_linked_to_previous = False
    cover_sec.header.paragraphs[0].clear()
    cover_sec.footer.is_linked_to_previous = False
    cover_sec.footer.paragraphs[0].clear()
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("Учреждение образования «Белорусский государственный университет транспорта»")
    for _ in range(5):
        blank_line(doc)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("ТЕХНИЧЕСКИЕ ТРЕБОВАНИЯ\nК РАЗРАБОТКЕ БИБЛИОТЕЧНОЙ ИНФОРМАЦИОННОЙ СИСТЕМЫ\nУЧЁТА ВЫДАЧИ И ВОЗВРАТА ИЗДАНИЙ\n(БИС)")
    for _ in range(4):
        blank_line(doc)
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("Лабораторная работа № 4\nТраньков Р.С., СИ-41")
    blank_line(doc)
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("На ")
    field(p, "NUMPAGES")
    p.add_run(" листах")
    for _ in range(4):
        blank_line(doc)
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("Гомель — 2026")

    body_sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    body_sec.top_margin = Cm(2.8)
    set_report_header(body_sec, doc)
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("СОДЕРЖАНИЕ").bold = True
    blank_line(doc)
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Cm(0)
    field(p, 'TOC \\o "1-2" \\h \\z \\u')
    doc.add_page_break()

    heading(doc, "1", "Введение")
    normal_paragraph(doc, "Документ задаёт проверяемые требования к библиотечной информационной системе, которая помогает библиотекарю учитывать выдачу и возврат изданий. Основа требований — IDEF0-модель варианта 18 из предыдущей лабораторной работы. Вход в систему и разграничение прав уточняют механизм M2 этой модели.")
    normal_paragraph(doc, "Формулировка задания о «зашифрованном виде» сохранена на вводной странице. В проектном решении пароль хранится как хеш Argon2id с солью, а клиентская проверка дублируется сервером, чтобы её нельзя было обойти прямым запросом. Оба уточнения требуют согласования. Полный текст СТП 09150.11.118-2009 в исходных материалах не представлен, поэтому соответствие всем его пунктам здесь не заявляется.")

    heading(doc, "2", "Общие сведения")
    normal_paragraph(doc, "Полное название: библиотечная информационная система учёта выдачи и возврата изданий (БИС). Пользователи системы — библиотекарь и администратор. Читатель участвует в процессе выдачи и возврата, но доступ читателя к самой БИС в исходной модели не предусмотрен.")
    normal_paragraph(doc, "Исходные материалы: IDEF0-модель «Учитывать выдачу и возврат изданий», словарь дуг ICOM и задание лабораторной работы № 4. Каталог и учёт читателей входят в БИС, а не считаются внешними системами. Смежный почтовый сервер для уведомлений добавлен как проектное допущение; его интерфейс описан ниже.")
    normal_paragraph(doc, "Числовые значения, отсутствующие в задании, приняты для учебного проекта и должны быть подтверждены библиотекой до реализации. Они перечислены в таблице 2.1.")
    blank_line(doc)
    table(doc, "Таблица 2.1 – Проектные допущения", ["Параметр", "Принятое значение", "Основание проверки"], [
        ["Выдача", "Не более 5 открытых выдач; срок 14 календарных дней", "Граничные сценарии F-02"],
        ["Копии", "Ежедневно; хранение 30 дней; восстановление до 4 часов", "Испытание Q-02"],
        ["Нагрузка", "20 одновременных сотрудников; 95 % поисков до 3 секунд", "Нагрузочное испытание Q-03"],
        ["Доступность", "Не менее 99 % за календарный месяц", "Мониторинг Q-04"],
        ["Журнал", "Хранение событий доступа 90 дней", "Проверка A-07"],
        ["Пароль", "8–128 символов; Argon2id: 19 МиБ, t=2, p=1", "Граничные тесты A-02, A-03"],
        ["Защита входа", "5 ошибок — блокировка 15 минут", "Проверка A-08"],
        ["Восстановление", "Код 15 минут; 5 попыток", "Проверка A-06"],
        ["Браузеры", "Chrome, Edge, Firefox 128+", "Проверка S-02"],
        ["Передача данных", "TLS 1.2 и выше", "Проверка I-01"],
    ], [4.0, 8.0, 5.0])
    blank_line(doc)
    normal_paragraph(doc, "Внешний почтовый сервер принимает уведомления от БИС по SMTP с TLS. БИС передаёт адрес, тему и текст сообщения, получает результат доставки и повторяет отправку после сбоя. Иных внешних интерфейсов исходная модель не задаёт; добавлять их без описания источника данных нельзя.")

    heading(doc, "3", "Назначение и цели создания")
    heading(doc, "3.1", "Назначение", level=2)
    normal_paragraph(doc, "БИС нужна для поиска экземпляров, проверки условий выдачи, оформления выдачи и возврата, контроля сроков, уведомлений и отчётов. Система не ведёт закупки фонда и финансовый учёт штрафов: эти процессы лежат за границами IDEF0-модели.")
    heading(doc, "3.2", "Цели создания", level=2)
    normal_paragraph(doc, "Система должна заменить разрозненные записи единым журналом операций, поддерживать актуальный статус каждого экземпляра и показывать библиотекарю ограничения до оформления выдачи. Доступ к операциям получают только сотрудники с назначенными правами.")

    heading(doc, "4", "Характеристика объекта автоматизации")
    normal_paragraph(doc, "Контекстная функция A-0 получает запрос читателя (I1) либо возвращённое издание (I2). Работой управляют правила библиотеки (C1), каталог и учётные правила (C2), сроки и ограничения (C3). Результатами служат выданное издание (O1), обновлённый статус фонда (O2), уведомление (O3) и отчёт (O4). Библиотекарь (M1) работает через БИС (M2) с библиотечным фондом (M3). Состав функций показан в таблице 4.1.")
    blank_line(doc)
    table(doc, "Таблица 4.1 – Функции и данные модели", ["Функция", "Содержание", "Связанные данные"], [
        ["A1", "Проверить читателя и доступность издания", "I1; C1–C3; M1–M3 → W1"],
        ["A2", "Зарегистрировать выдачу издания", "W1; C1–C3; M1–M3 → W2, O1"],
        ["A3", "Зарегистрировать возврат и контролировать срок", "I2, W2; C1–C3; M1–M3 → O2, W3"],
        ["A4", "Сформировать уведомления и отчёты", "W3; C1; M2 → O3, O4"],
    ], [2.0, 8.1, 6.9])
    blank_line(doc)
    normal_paragraph(doc, "Состав и формат каждой дуги ICOM приведены в таблице 4.2. W1 означает решение по заявке, W2 — запись выдачи, W3 — сведения о просрочке. Вход в БИС и роль администратора не показаны в исходной модели обслуживания читателя. Требования к ним являются проектным расширением механизма M2 и прослеживаются через диаграмму Б.1, а не выдаются за функцию A1–A4.")
    blank_line(doc)
    table(doc, "Таблица 4.2 – Дуги ICOM и состав данных", ["Код", "Назначение", "Состав или форма представления"], [
        ["I1", "Запрос читателя", "Идентификаторы читателя и экземпляра; запись в форме заявки"],
        ["I2", "Возвращённое издание", "Идентификатор экземпляра и состояние при приёмке"],
        ["C1", "Правила библиотеки", "Локальный регламент обслуживания и роли"],
        ["C2", "Каталог фонда", "Карточки изданий, экземпляров и их статусы"],
        ["C3", "Сроки и ограничения", "Проектные значения: 14 дней, 5 открытых выдач"],
        ["O1", "Выданное издание", "Экземпляр и подтверждённая запись выдачи"],
        ["O2", "Статус фонда", "Идентификатор экземпляра, состояние и доступность"],
        ["O3", "Уведомление", "Адрес, тема, текст, результат доставки"],
        ["O4", "Отчёт", "PDF и CSV; период, операции, итоговые количества"],
        ["M1–M3", "Механизмы", "Библиотекарь, БИС и библиотечный фонд"],
        ["W1", "Решение по заявке", "Читатель, экземпляр, принято/отказ, причина"],
        ["W2", "Запись выдачи", "Читатель, экземпляр, даты выдачи и возврата"],
        ["W3", "Сведения о просрочке", "Номер выдачи, дата срока и число дней просрочки"],
    ], [2.0, 5.0, 10.0])
    blank_line(doc)

    heading(doc, "5", "Требования к системе")
    heading(doc, "5.1", "Требования к структуре", level=2)
    add_requirements(doc, ["S-01", "S-02"])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(STRUCTURE_DIAGRAM), width=Cm(12.5))
    p = doc.add_paragraph("Рисунок 5.1 – Структура БИС")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    blank_line(doc)
    heading(doc, "5.2", "Требования к процедурам доступа", level=2)
    add_requirements(doc, ["A-01", "A-02", "A-03", "A-04", "A-05", "A-06", "A-07", "A-08", "A-09"])
    normal_paragraph(doc, "Матрица прав: библиотекарь ищет, регистрирует читателей, выдаёт и принимает экземпляры, формирует отчёты. Администратор создаёт учётные записи сотрудников, меняет роли, запускает резервное копирование и восстановление; выдачу и возврат он не оформляет.")
    heading(doc, "5.3", "Функциональные требования", level=2)
    add_requirements(doc, ["F-01", "F-02", "F-03", "F-04", "F-05", "F-06", "F-07"])
    heading(doc, "5.4", "Требования к информационной безопасности", level=2)
    add_requirements(doc, ["I-01", "I-02"])
    heading(doc, "5.5", "Требования к патентной чистоте", level=2)
    add_requirements(doc, ["PAT-01"])
    heading(doc, "5.6", "Надёжность и производительность", level=2)
    add_requirements(doc, ["Q-01", "Q-02", "Q-03", "Q-04"])

    heading(doc, "6", "Требования к видам обеспечения")
    heading(doc, "6.1", "Информационное обеспечение", level=2)
    add_requirements(doc, ["D-01", "D-02"])
    normal_paragraph(doc, "Каталог и карточки экземпляров задают доступность фонда. Запись выдачи связывает читателя, экземпляр, дату выдачи, срок и факт возврата. События доступа хранятся отдельно от данных библиотечных операций.")
    heading(doc, "6.2", "Математическое обеспечение", level=2)
    normal_paragraph(doc, "Просрочка определяется сравнением текущей даты с датой возврата в открытой записи. Лимит выдачи и срок пользования приведены в таблице 2.1. Штрафы не рассчитываются: финансовый учёт находится вне границ модели.")
    heading(doc, "6.3", "Программное обеспечение", level=2)
    normal_paragraph(doc, "БИС включает браузерный интерфейс, серверную часть и хранилище данных. Сервер проверяет роли по A-04, а пароли хранятся по A-03. Смежный почтовый сервер используется только для исходящих уведомлений.")
    add_requirements(doc, ["INT-01"])
    heading(doc, "6.4", "Техническое обеспечение", level=2)
    normal_paragraph(doc, "Поддерживаемые браузеры определены в S-02. Нагрузочный ориентир и критерий ответа системы заданы в Q-03; фактическую ёмкость сервера подбирают по результату нагрузочного испытания.")
    heading(doc, "6.5", "Организационное обеспечение", level=2)
    normal_paragraph(doc, "Ответственность библиотекаря и администратора разделена матрицей ролей в 5.2. Порядок предоставления и отзыва доступа утверждается библиотекой до ввода БИС в работу.")

    heading(doc, "7", "Состав и содержание работ")
    normal_paragraph(doc, "Этапы разработки и ожидаемые результаты приведены в таблице 7.1.")
    blank_line(doc)
    table(doc, "Таблица 7.1 – Этапы разработки системы", ["Этап", "Результат", "Ответственный и срок"], [
        ["Уточнение требований", "Согласованные роли, ограничения и критерии приёмки.", "Заказчик и аналитик; 2 недели"],
        ["Проектирование", "Схема данных, интерфейсы и модель доступа.", "Разработчик; 2 недели"],
        ["Разработка", "Каталог, выдача, возврат, доступ, уведомления и отчёты.", "Разработчик; 6 недель"],
        ["Проверка", "Протоколы испытаний по ID требований.", "Тестировщик; 2 недели"],
        ["Передача", "Документация, резервная копия и обучение.", "Заказчик и разработчик; 1 неделя"],
    ], [3.4, 8.6, 5.0])
    blank_line(doc)
    normal_paragraph(doc, "Указанные сроки и состав исполнителей являются проектным планом, а не утверждённым графиком заказчика. Дата начала работ в задании не дана.")

    heading(doc, "8", "Порядок приёмки в эксплуатацию")
    normal_paragraph(doc, "Приёмка проводится по сценариям, указанным у каждого требования. Для бизнес-функций используются тестовые читатели и экземпляры; для доступа — учётные записи библиотекаря и администратора. Испытания проходят на отдельной копии данных. Основные проверки сведены в таблицу 8.1.")
    blank_line(doc)
    table(doc, "Таблица 8.1 – Основные приёмочные проверки", ["Группа", "Основная проверка", "Результат"], [
        ["Структура S-01–S-02", "Функции подсистем в трёх браузерах", "Все подсистемы доступны согласно роли."],
        ["Доступ A-01–A-09", "Вход, роли, политика пароля, сброс, тайм-аут и блокировка", "Недопустимые действия отклоняются по каждому ID."],
        ["Операции F-01–F-07", "Каталог, читатели, выдача, возврат, уведомления, отчёт", "Записи и статус соответствуют операциям."],
        ["Качество Q-01–Q-04", "Параллельная выдача, копия, нагрузка и доступность", "Достигнуты значения из таблицы 2.1."],
        ["Данные D-01–D-02", "Поля запроса и отчёта, связи между записями", "Обязательные поля и связи проверены."],
        ["Защита I-01–I-02", "TLS и минимальный набор полей по ролям", "Незащищённые запросы и лишние поля отсутствуют."],
        ["Прочие INT-01, PAT-01, DOC-01", "Почта, перечень лицензий, комплект документов", "Результаты представлены в протоколе приёмки."],
    ], [3.0, 7.0, 7.0])
    blank_line(doc)
    normal_paragraph(doc, "При обнаружении ошибки фиксируются требование, входные данные, ожидаемый и фактический результат. Повторное испытание проводится после исправления.")

    heading(doc, "9", "Требования к документированию")
    add_requirements(doc, ["DOC-01"])
    normal_paragraph(doc, "Текст технических требований и матрица прослеживаемости передаются в редактируемом формате. Исходная IDEF0-модель представлена на рисунках А.1–А.5 в приложении А, схемы процессов доступа — на рисунках Б.1 и Б.2 в приложении Б. Связь требований с моделью приведена в приложении В.")
    normal_paragraph(doc, "При подготовке требований к паролям и сеансам учтены рекомендации OWASP по хранению паролей [1] и управлению сеансами [2]. Исходное задание сохраняет приоритет для параметра бездействия: 10 минут.")

    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    blank_line(doc)
    normal_paragraph(doc, "1\u00a0OWASP Cheat Sheet Series. Password Storage Cheat Sheet : [сайт]. – URL: https://cheatsheetseries.owasp.org/cheatsheets/Password_Storage_Cheat_Sheet.html (дата обращения: 06.10.2026).")
    normal_paragraph(doc, "2\u00a0OWASP Cheat Sheet Series. Session Management Cheat Sheet : [сайт]. – URL: https://cheatsheetseries.owasp.org/cheatsheets/Session_Management_Cheat_Sheet.html (дата обращения: 06.10.2026).")
    normal_paragraph(doc, "3\u00a0Лабораторная работа № 4. Разработка технических требований : задание и пример оформления. – Учебный материал, предоставленный преподавателем.")
    normal_paragraph(doc, "4\u00a0СТП 09150.11.118-2009 : документ, указанный в задании. Полный текст при подготовке работы не предоставлен.")

    idef0_sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    idef0_sec.orientation = WD_ORIENT.PORTRAIT
    idef0_sec.page_width = Cm(21)
    idef0_sec.page_height = Cm(29.7)
    idef0_sec.left_margin = Cm(3)
    idef0_sec.right_margin = Cm(1)
    idef0_sec.top_margin = Cm(2.8)
    idef0_sec.bottom_margin = Cm(2.5)
    set_report_header(idef0_sec, doc)
    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("ПРИЛОЖЕНИЕ А\n(справочное)\nИсходная модель IDEF0")
    blank_line(doc)
    model_diagrams = [
        ("A-0_контекстная_диаграмма.png", "Рисунок А.1. Контекстная диаграмма A-0"),
        ("A0_декомпозиция.png", "Рисунок А.2. Декомпозиция главной функции A0"),
        ("A1_проверка_заявки.png", "Рисунок А.3. Проверка заявки A1"),
        ("A2_выдача_издания.png", "Рисунок А.4. Выдача издания A2"),
        ("A3_возврат_и_контроль.png", "Рисунок А.5. Возврат и контроль срока A3"),
    ]
    for index, (filename, caption) in enumerate(model_diagrams):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        if index > 0:
            p.paragraph_format.page_break_before = True
        image_width = Cm(17)
        p.add_run().add_picture(str(IDEF0_EXPORTS / filename), width=image_width)
        blank_line(doc).paragraph_format.keep_with_next = True
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_together = True
        r = p.add_run(caption.replace(". ", " – ", 1))
        r.bold = True
        r.font.size = Pt(12)

    portrait_sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    portrait_sec.orientation = WD_ORIENT.PORTRAIT
    portrait_sec.page_width = Cm(21)
    portrait_sec.page_height = Cm(29.7)
    portrait_sec.left_margin = Cm(3)
    portrait_sec.right_margin = Cm(1)
    portrait_sec.top_margin = Cm(2.8)
    portrait_sec.bottom_margin = Cm(2.5)
    set_report_header(portrait_sec, doc)

    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("ПРИЛОЖЕНИЕ Б\n(справочное)\nПроцессы доступа")
    blank_line(doc)
    normal_paragraph(doc, "На рисунке Б.1 показаны вход, повторная попытка, проверка роли, цикл операций и завершение сеанса. На рисунке Б.2 показаны создание и восстановление пароля. Ответвления помечены условиями в квадратных скобках.")
    blank_line(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(DIAGRAM), width=Cm(14.2))
    blank_line(doc).paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("Рисунок Б.1 – Процесс аутентификации и авторизации")
    r.bold = True
    r.font.size = Pt(12)

    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(RESET_DIAGRAM), width=Cm(15.2))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("Рисунок Б.2 – Создание пароля и восстановление доступа")
    r.bold = True
    r.font.size = Pt(12)

    doc.add_page_break()
    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("ПРИЛОЖЕНИЕ В\n(справочное)\nПрослеживаемость требований")
    blank_line(doc)
    normal_paragraph(doc, "В графе «Источник» указаны элементы исходной IDEF0-модели или проектные дополнения. На диаграммах A означает функцию, I — вход, C — управление, O — выход, M — механизм, W — внутреннюю дугу. В идентификаторах требований префиксы S, A, F, Q, I, D, INT, PAT и DOC обозначают соответственно структуру, доступ, функции, качество, безопасность, данные, интерфейс, патентную чистоту и документы; дефис отделяет префикс от номера. Связь требований с моделью приведена в таблице В.1.")
    blank_line(doc)
    table(doc, "Таблица В.1 – Прослеживаемость требований", ["ID", "Источник", "Смысл связи"], [
        [req_id, source, TRACE_HINTS.get(req_id, "Проектное дополнение к модели и заданию.")]
        for req_id, statement, _, source in REQUIREMENTS
    ], [2.0, 4.2, 10.8])
    doc.core_properties.title = "Лабораторная работа № 4. Технические требования к библиотечной информационной системе"
    doc.core_properties.author = "Роман Траньков"
    doc.save(REPORT)


if __name__ == "__main__":
    make_revised_diagrams()
    build_document()
    print(REPORT)
